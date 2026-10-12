"""Shared parsing and row-by-row processing for admin bulk imports."""

import csv
from dataclasses import dataclass
from io import BytesIO, StringIO
from pathlib import Path
from typing import Callable
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from fastapi import HTTPException, UploadFile, status
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError
from sqlalchemy.exc import IntegrityError
from sqlmodel import Session


MAX_IMPORT_ROWS = 500
MAX_IMPORT_FILE_BYTES = 5 * 1024 * 1024


@dataclass(frozen=True)
class ImportRow:
    number: int
    data: object


class BulkRowError(ValueError):
    """A row-specific import error suitable for returning to the caller."""


def json_import_rows(rows: list[object]) -> list[ImportRow]:
    if len(rows) > MAX_IMPORT_ROWS:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"Bulk imports are limited to {MAX_IMPORT_ROWS} rows",
        )
    return [ImportRow(number=index, data=row) for index, row in enumerate(rows, start=1)]


def uploaded_import_rows(upload: UploadFile) -> list[ImportRow]:
    filename = upload.filename or ""
    extension = Path(filename).suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Unsupported file type; upload a .csv or .xlsx file",
        )

    content = upload.file.read(MAX_IMPORT_FILE_BYTES + 1)
    if len(content) > MAX_IMPORT_FILE_BYTES:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail=f"Import files are limited to {MAX_IMPORT_FILE_BYTES // (1024 * 1024)} MB",
        )
    if not content:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Import file is empty")

    if extension == ".csv":
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="CSV file must use UTF-8 encoding",
            ) from exc
        return _parse_csv(text)
    return _parse_xlsx(content)


def _validated_headers(values: tuple[object, ...] | list[object] | None) -> list[str]:
    if not values:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Import file must have a header row",
        )
    headers = [str(value).strip() if value is not None else "" for value in values]
    if any(not header for header in headers):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Import file headers cannot be blank",
        )
    if len({header.casefold() for header in headers}) != len(headers):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Import file headers must be unique",
        )
    return headers


def _parse_csv(text: str) -> list[ImportRow]:
    reader = csv.reader(StringIO(text))
    headers = _validated_headers(next(reader, None))
    rows: list[ImportRow] = []
    for values in reader:
        row_number = reader.line_num
        if not values or not any(value.strip() for value in values):
            continue
        if len(rows) >= MAX_IMPORT_ROWS:
            raise HTTPException(
                status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                detail=f"Bulk imports are limited to {MAX_IMPORT_ROWS} rows",
            )
        row_data: dict[str, object] = {
            header: _normalize_cell(value) for header, value in zip(headers, values)
        }
        if len(values) > len(headers):
            row_data["__row_error__"] = "Row has more values than the header"
        rows.append(ImportRow(number=row_number, data=row_data))
    return rows


def _parse_xlsx(content: bytes) -> list[ImportRow]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ParseError, ValueError) as exc:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Could not read XLSX file; ensure it is a valid Excel workbook",
        ) from exc

    try:
        worksheet = workbook.active
        rows_iterator = worksheet.iter_rows(values_only=True)
        headers = _validated_headers(next(rows_iterator, None))
        rows: list[ImportRow] = []
        for row_number, values in enumerate(rows_iterator, start=2):
            if not values or not any(value is not None and str(value).strip() for value in values):
                continue
            if len(rows) >= MAX_IMPORT_ROWS:
                raise HTTPException(
                    status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
                    detail=f"Bulk imports are limited to {MAX_IMPORT_ROWS} rows",
                )
            row_data: dict[str, object] = {
                header: _normalize_cell(value) for header, value in zip(headers, values)
            }
            if len(values) > len(headers) and any(value is not None for value in values[len(headers):]):
                row_data["__row_error__"] = "Row has more values than the header"
            rows.append(ImportRow(number=row_number, data=row_data))
        return rows
    finally:
        workbook.close()


def _normalize_cell(value: object) -> object:
    if isinstance(value, str) and not value.strip():
        return None
    return value


def process_import_rows(
    rows: list[ImportRow],
    db: Session,
    create_record: Callable[[dict[str, object]], object],
) -> dict[str, object]:
    created_records: list[dict[str, object]] = []
    errors: list[dict[str, object]] = []

    for row in rows:
        try:
            if not isinstance(row.data, dict) or any(not isinstance(key, str) for key in row.data):
                raise BulkRowError("Each row must be an object with named fields")
            if "__row_error__" in row.data:
                raise BulkRowError(str(row.data["__row_error__"]))
            with db.begin_nested():
                record = create_record(row.data)
                db.add(record)
                db.flush()
            record_id = getattr(record, "id", None)
            if record_id is not None:
                created_records.append({"row": row.number, "id": str(record_id)})
        except ValidationError as exc:
            messages = [
                f"{'.'.join(str(part) for part in error['loc'])}: {error['msg']}"
                for error in exc.errors()
            ]
            errors.append({"row": row.number, "message": "; ".join(messages)})
        except (BulkRowError, IntegrityError) as exc:
            message = str(exc) if isinstance(exc, BulkRowError) else "Database constraint violation"
            errors.append({"row": row.number, "message": message})

    db.commit()
    return {
        "total_rows": len(rows),
        "created": len(created_records),
        "failed": len(errors),
        "created_records": created_records,
        "errors": errors,
    }
